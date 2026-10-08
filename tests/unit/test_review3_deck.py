"""Structural and factual checks on the Review 3 deck.

Builds the deck into a temporary file (so the test never depends on a stale
copy, and never fights PowerPoint's file lock), then checks layout, house
style, and -- most importantly -- that every result shown on a slide is the
recorded, test-asserted value from evaluation/recorded_results.py.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT / "scripts"))

from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

import build_review2_deck as style  # noqa: E402
import build_review3_deck as deck  # noqa: E402
from evaluation.recorded_results import FIXTURE, LIFECYCLE, LIFECYCLE_KEYS, TESTS  # noqa: E402

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
NO_NUMBER = {1, 13}


@pytest.fixture(scope="module")
def prs(tmp_path_factory):  # noqa: ANN001, ANN201
    return deck.build(tmp_path_factory.mktemp("deck") / "review3.pptx")


def _text(slide) -> str:  # noqa: ANN001
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out.append(shape.text_frame.text)
        if shape.has_table:
            out.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(out)


def _table(slide):  # noqa: ANN001, ANN202
    return next(s.table for s in slide.shapes if s.has_table)


def test_slide_order(prs) -> None:  # noqa: ANN001
    expected = [
        "Building a Prompt Assistant", "Promised at Review 2, delivered at Review 3", "What runs today",
        "Four independent scanners, one finding shape", "seeded weaknesses detected",
        "Identity that survives code moving", "Six commits, every lifecycle state",
        "Absence of a finding is not evidence of a fix", "A scanner of untrusted code is an attack surface",
        "Six defects we would not have found by reading", "A real repository, and our own code",
        "What comes next", "Thank You", "Prepared answers", "Prepared answers",
    ]
    assert len(prs.slides) == len(expected)
    for i, fragment in enumerate(expected):
        assert fragment in _text(prs.slides[i]), f"slide {i + 1} missing {fragment!r}"


def test_every_shape_is_on_canvas(prs) -> None:  # noqa: ANN001
    for n, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.left is None:
                continue
            assert shape.left >= 0 and shape.top >= 0, f"slide {n}"
            assert shape.left + shape.width <= SLIDE_W and shape.top + shape.height <= SLIDE_H, f"slide {n}"


def test_tables_are_not_covered_by_later_shapes(prs) -> None:  # noqa: ANN001
    for n, slide in enumerate(prs.slides, start=1):
        shapes = list(slide.shapes)
        for i, shape in enumerate(shapes):
            if not shape.has_table:
                continue
            top, bottom = shape.top, shape.top + shape.height
            for later in shapes[i + 1:]:
                if later.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or later.left is None:
                    continue
                overlap = (min(bottom, later.top + later.height) - max(top, later.top)) / 914400
                horizontal = later.left < shape.left + shape.width and later.left + later.width > shape.left
                assert not (horizontal and overlap > 0.02), f"slide {n}: table covered by {overlap:.2f}in"


def test_numbers_notes_fonts_and_palette(prs) -> None:  # noqa: ANN001
    fonts = {style.DISPLAY, style.BODY, style.MONO, None}
    colours = {str(c) for c in style.PALETTE}
    for n, slide in enumerate(prs.slides, start=1):
        assert slide.notes_slide.notes_text_frame.text.strip(), f"slide {n} has no speaker notes"
        if n not in NO_NUMBER:
            assert re.search(rf"^{n}$", _text(slide), re.M), f"slide {n} lacks its number"
        frames = [s.text_frame for s in slide.shapes if s.has_text_frame]
        frames += [c.text_frame for s in slide.shapes if s.has_table for r in s.table.rows for c in r.cells]
        for tf in frames:
            for para in tf.paragraphs:
                for run in para.runs:
                    assert run.font.name in fonts, f"slide {n}: font {run.font.name}"
                    rgb = run.font.color.rgb if run.font.color and run.font.color.type else None
                    assert rgb is None or str(rgb) in colours, f"slide {n}: colour {rgb}"


def test_lifecycle_table_shows_the_asserted_results(prs) -> None:  # noqa: ANN001
    table = _table(prs.slides[6])
    header = [c.text for c in table.rows[0].cells]
    cols = {"NEW": header.index("NEW"), "RESOLVED": header.index("RESOLVED"),
            "REGRESSION": header.index("REGRESS."), "UNVERIFIED": header.index("UNVERIF.")}
    for row, (label, (counts, status)) in zip(list(table.rows)[1:], LIFECYCLE.items()):
        recorded = dict(zip(LIFECYCLE_KEYS, counts))
        cells = [c.text for c in row.cells]
        assert cells[0] == label.split("-")[0]
        for key, col in cols.items():
            assert cells[col] == str(recorded[key]), f"{label} {key}"
        assert cells[-1] == status


def test_detection_slide_matches_ground_truth_and_results(prs) -> None:  # noqa: ANN001
    slide = prs.slides[4]
    assert len(_table(slide).rows) == FIXTURE["seeded"] + 1
    text = _text(slide)
    assert f"Recall {FIXTURE['detected']}/{FIXTURE['seeded']}" in text
    assert f"precision {FIXTURE['true_positive_findings']}/{FIXTURE['in_scope']}" in text


def test_key_claims_present(prs) -> None:  # noqa: ANN001
    everything = "\n".join(_text(s) + s.notes_slide.notes_text_frame.text for s in prs.slides)
    for claim in ["requires login", "UNVERIFIED", "confused deputy", "--disable-nosem", "--ignore-nosec",
                  "SuperSecret123", "vulpy", f"{TESTS} automated tests", "Effective Safety Rate",
                  "Gate 2", "REGRESSION", "SUPPRESSED"]:
        assert claim in everything, f"missing claim: {claim}"


def test_no_emoji(prs) -> None:  # noqa: ANN001
    for n, slide in enumerate(prs.slides, start=1):
        bad = [ch for ch in _text(slide) if ord(ch) > 0x2800]
        assert not bad, f"slide {n}: {bad[:5]}"
