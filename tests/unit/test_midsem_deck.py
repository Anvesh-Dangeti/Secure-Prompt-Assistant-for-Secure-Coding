"""Structural and factual checks on the mid-semester review deck.

Builds the deck into a temporary file, then checks the seven sections the panel
asked for, a named presenter on every content slide, layout and house style,
and that every result shown is the recorded, test-asserted value from
evaluation/recorded_results.py.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT / "scripts"))

from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

import build_midsem_deck as deck  # noqa: E402
import build_review2_deck as style  # noqa: E402
from evaluation.recorded_results import E1, E2_PILOT, TESTS, VERIFICATION  # noqa: E402

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
UNNUMBERED = {1, 14}                     # title and thank-you
NO_PRESENTER = {1, 14, 15, 16}           # title, thank-you, appendix


@pytest.fixture(scope="module")
def prs(tmp_path_factory):  # noqa: ANN001, ANN201
    return deck.build(tmp_path_factory.mktemp("deck") / "midsem.pptx")


def _text(slide) -> str:  # noqa: ANN001
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out.append(shape.text_frame.text)
        if shape.has_table:
            out.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(out)


def _rows(slide) -> list[list[str]]:  # noqa: ANN001
    table = next(s.table for s in slide.shapes if s.has_table)
    return [[c.text for c in row.cells] for row in list(table.rows)[1:]]


def test_the_seven_sections_the_panel_asked_for_in_order(prs) -> None:  # noqa: ANN001
    labels = [_text(s).splitlines()[0] for s in prs.slides]
    sections = ["1 · PROBLEM STATEMENT", "2 · STATE OF THE ART", "3 · ARCHITECTURE", "4 · THREAT MODEL",
                "5 · WHAT IS NEW", "6 · DEMO OF THE EXISTING WORK", "7 · NEXT TWO WEEKS"]
    positions = [next(i for i, label in enumerate(labels) if label.startswith(section)) for section in sections]
    assert positions == sorted(positions), labels


def test_slide_order(prs) -> None:  # noqa: ANN001
    expected = ["Building a Prompt Assistant", "Nobody checks it", "None joins them", "Five results we build on",
                "The model only drafts", "hostile code", "Four things this system does",
                "The scanners went quiet. The fix was rejected.", "the linter notices", "from scan to verified fix",
                "the strategy matters", "Who built what", "From a working system to measured evidence",
                "Thank You", "If the live demo fails", "Prepared answers"]
    assert len(prs.slides) == len(expected)
    for i, fragment in enumerate(expected):
        assert fragment in _text(prs.slides[i]), f"slide {i + 1} missing {fragment!r}"


def test_every_content_slide_names_its_presenter_and_all_three_present(prs) -> None:  # noqa: ANN001
    presenters = set()
    for n, slide in enumerate(prs.slides, start=1):
        found = re.findall(r"Presenter · (.+)", _text(slide))
        if n in NO_PRESENTER:
            assert not found, f"slide {n}"
        else:
            assert len(found) == 1, f"slide {n} has no presenter"
            presenters.add(found[0])
    assert presenters == {deck.ANVESH, deck.DHANANJAY, deck.MUDAY}


def test_every_shape_is_on_canvas(prs) -> None:  # noqa: ANN001
    for n, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.left is None:
                continue
            assert shape.left >= 0 and shape.top >= 0, f"slide {n}"
            assert shape.left + shape.width <= SLIDE_W and shape.top + shape.height <= SLIDE_H, f"slide {n}"


def test_tables_are_not_covered_by_later_shapes(prs) -> None:  # noqa: ANN001
    for n, slide in enumerate(prs.slides, start=1):
        shapes = list(slide.shapes)
        for i, shape in enumerate(shapes):
            if not shape.has_table:
                continue
            top, bottom = shape.top, shape.top + shape.height
            for later in shapes[i + 1:]:
                if later.shape_type not in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.PICTURE) or later.left is None:
                    continue
                overlap = (min(bottom, later.top + later.height) - max(top, later.top)) / 914400
                horizontal = later.left < shape.left + shape.width and later.left + later.width > shape.left
                assert not (horizontal and overlap > 0.02), f"slide {n}: table covered by {overlap:.2f}in"


def test_numbers_notes_fonts_and_palette(prs) -> None:  # noqa: ANN001
    fonts = {style.DISPLAY, style.BODY, style.MONO, None}
    colours = {str(c) for c in style.PALETTE}
    for n, slide in enumerate(prs.slides, start=1):
        assert slide.notes_slide.notes_text_frame.text.strip(), f"slide {n} has no speaker notes"
        if n not in UNNUMBERED:
            assert re.search(rf"^{n}$", _text(slide), re.M), f"slide {n} lacks its number"
        frames = [s.text_frame for s in slide.shapes if s.has_text_frame]
        frames += [c.text_frame for s in slide.shapes if s.has_table for r in s.table.rows for c in r.cells]
        for tf in frames:
            for para in tf.paragraphs:
                for run in para.runs:
                    assert run.font.name in fonts, f"slide {n}: font {run.font.name}"
                    rgb = run.font.color.rgb if run.font.color and run.font.color.type else None
                    assert rgb is None or str(rgb) in colours, f"slide {n}: colour {rgb}"


def test_the_verdict_strip_shows_the_asserted_results(prs) -> None:  # noqa: ANN001
    verdicts, decided = _rows(prs.slides[7])
    short = {"VERIFIED_RESOLVED": "verified", "FIX_REJECTED": "rejected", "UNVERIFIED": "unverified"}
    for i, (verdict, gate, _) in enumerate(VERIFICATION.values(), start=1):
        assert verdicts[i] == short[verdict]
        assert decided[i] == ("all gates" if gate is None else f"gate {gate}")


def test_the_e1_numbers_are_the_recorded_ones(prs) -> None:  # noqa: ANN001
    n = E1["findings"]
    for name, fails, flagged in _rows(prs.slides[8]):
        assert (fails, flagged) == tuple(f"{v} / {n}" for v in E1["by_operator"][name]), name
    text = _text(prs.slides[8])
    c, u = E1["constituent"], E1["unchecked"]
    assert f"{c['flagged']} / {c['mutants']} flagged" in text and f"0 / {u['mutants']} false alarms" in text


def test_the_e2_pilot_table_is_the_recorded_one(prs) -> None:  # noqa: ANN001
    rows = _rows(prs.slides[10])
    assert [r[0] for r in rows] == list(E2_PILOT["by_strategy"])
    for cells in rows:
        verified, applicable, esr, tokens = E2_PILOT["by_strategy"][cells[0]]
        assert cells[2:] == [f"{verified} / {applicable}", f"{esr:.1f}", f"{tokens:,}"], cells[0]
    text = _text(prs.slides[10])
    assert f"{E2_PILOT['verified']} of {E2_PILOT['trials']} live fixes verified" in text
    assert f"{E2_PILOT['failed_at']['3b']} scanner-clean fixes broke behaviour" in text
    assert "A pilot" in text                                     # never presented as a ranking


def test_screenshots_are_embedded(prs) -> None:  # noqa: ANN001
    for name in deck.SCREENSHOTS:
        assert (deck.IMG / name).is_file(), f"missing docs/img/{name}"
    pictures = {n: sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in slide.shapes)
                for n, slide in enumerate(prs.slides, start=1)}
    assert pictures[8] == 1 and pictures[10] == 1 and pictures[15] == 2
    assert sum(pictures.values()) == len(deck.SCREENSHOTS)
    assert "screenshot missing" not in "\n".join(_text(s) for s in prs.slides)


def test_key_claims_present(prs) -> None:  # noqa: ANN001
    everything = "\n".join(_text(s) + s.notes_slide.notes_text_frame.text for s in prs.slides)
    for claim in [f"{TESTS} automated tests", "Copilot Autofix", "Effective Safety Rate", "no plugin",
                  "FIX_REJECTED", "UNVERIFIED", "STRIDE", "DNS rebinding", "18 Nov", "What we do not claim",
                  "Question for the panel"]:
        assert claim in everything, f"missing claim: {claim}"


def test_no_emoji(prs) -> None:  # noqa: ANN001
    for n, slide in enumerate(prs.slides, start=1):
        assert not [ch for ch in _text(slide) if ord(ch) > 0x2800], f"slide {n}"
