"""Structural checks on the generated Review 2 deck.

A deck that looks fine to the generator can still be broken when opened: shapes
off-canvas, a table with the wrong shape, a slide missing its notes, or a colour
that drifted off the house palette. These assertions are cheap and catch all of
that before the deck is in front of a review panel.

Run with the interpreter that has python-pptx (Python 3.12 on this machine):
    <py312> tests/unit/test_review2_deck.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
PROJECT_PARENT = REPO_ROOT.parent
DECK = PROJECT_PARENT / "Secure_Prompt_Assistant_Review2.pptx"
SLIDES_MD = REPO_ROOT / "docs" / "review2_slides.md"

sys.path.insert(0, str(REPO_ROOT / "scripts"))

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

import build_review2_deck as deck  # noqa: E402

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

ALLOWED_FONTS = {deck.DISPLAY, deck.BODY, deck.MONO, None}
ALLOWED_COLORS = {str(c) for c in deck.PALETTE}

# Slides with no page number by design (title and closing).
NO_NUMBER = {1, 13}


def _prs() -> Presentation:
    assert DECK.exists(), f"deck not built: {DECK}"
    return Presentation(str(DECK))


def _all_text(slide) -> str:
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out.append(shape.text_frame.text)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    out.append(cell.text)
    return "\n".join(out)


def test_slide_count_and_order() -> None:
    prs = _prs()
    assert len(prs.slides) == 16, f"expected 16 slides, got {len(prs.slides)}"

    expected = [
        "Building a Prompt Assistant",
        "Four questions, four answers",
        "The loop we proposed at Review 1 already ships",
        "Three empty columns",
        "Three papers, three findings we build on",
        "One language: Python",
        "8 core CWEs + 2 supporting",
        "What we now claim",
        "A fix that deletes the vulnerable code",
        "Where the changes land",
        "Five gaps, and how we measure against them",
        "Where we are",
        "Thank You",
        "Prepared answers",
        "Prepared answers",
        "Prepared answers",
    ]
    for i, fragment in enumerate(expected):
        text = _all_text(prs.slides[i])
        assert fragment in text, f"slide {i + 1}: missing headline fragment {fragment!r}"


def test_every_shape_is_on_canvas() -> None:
    prs = _prs()
    problems = []
    for n, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.left is None or shape.top is None:
                continue
            right = shape.left + (shape.width or 0)
            bottom = shape.top + (shape.height or 0)
            if shape.left < 0 or shape.top < 0 or right > SLIDE_W or bottom > SLIDE_H:
                problems.append(
                    f"slide {n}: {shape.shape_type} at "
                    f"({shape.left / 914400:.2f},{shape.top / 914400:.2f}) "
                    f"extends to ({right / 914400:.2f},{bottom / 914400:.2f})"
                )
    assert not problems, "shapes off canvas:\n  " + "\n  ".join(problems)


def test_tables_are_not_covered_by_later_shapes() -> None:
    """A table can be fully on-canvas and still be hidden by an opaque callout.

    This is the bug class that the on-canvas check misses: on slides 7 and 9 the
    bottom table row sat underneath the callout strip, so the last CWE and the
    whole of gate 3 were invisible when the deck was opened. Shapes render in
    z-order, so anything added after a table and overlapping it covers it.
    """
    prs = _prs()
    problems = []
    for n, slide in enumerate(prs.slides, start=1):
        shapes = list(slide.shapes)
        for i, shape in enumerate(shapes):
            if not shape.has_table:
                continue
            t_l, t_t = shape.left, shape.top
            t_r, t_b = t_l + shape.width, t_t + shape.height
            for later in shapes[i + 1:]:
                if later.left is None:
                    continue
                # Text boxes are transparent; only filled autoshapes occlude.
                if later.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
                    continue
                l_l, l_t = later.left, later.top
                l_r, l_b = l_l + (later.width or 0), l_t + (later.height or 0)
                if l_l < t_r and l_r > t_l and l_t < t_b and l_b > t_t:
                    overlap_in = (min(t_b, l_b) - max(t_t, l_t)) / 914400
                    if overlap_in > 0.02:
                        problems.append(
                            f"slide {n}: shape at y={l_t / 914400:.2f}in covers "
                            f"{overlap_in:.2f}in of the table "
                            f"(table spans {t_t / 914400:.2f}-{t_b / 914400:.2f}in)"
                        )
    assert not problems, "content hidden behind later shapes:\n  " + "\n  ".join(problems)


def test_dense_tables_have_expected_shape() -> None:
    prs = _prs()

    def table_on(index):
        for shape in prs.slides[index].shapes:
            if shape.has_table:
                return shape.table
        raise AssertionError(f"slide {index + 1} has no table")

    matrix = table_on(3)  # slide 4 — capability matrix
    assert (len(matrix.rows), len(matrix.columns)) == (8, 8), (
        f"capability matrix is {len(matrix.rows)}x{len(matrix.columns)}, expected 8x8"
    )

    cwes = table_on(6)  # slide 7 — CWE scope
    assert (len(cwes.rows), len(cwes.columns)) == (11, 4), (
        f"CWE table is {len(cwes.rows)}x{len(cwes.columns)}, expected 11x4"
    )

    # Column widths must sum to the content width, or the table overhangs.
    for name, table, expected in (("matrix", matrix, 12.10), ("cwes", cwes, 12.10)):
        total = sum(c.width for c in table.columns) / 914400
        assert abs(total - expected) < 0.02, f"{name} columns sum to {total:.2f}in"


def test_slide_numbers_and_notes_present() -> None:
    prs = _prs()
    for n, slide in enumerate(prs.slides, start=1):
        notes = slide.notes_slide.notes_text_frame.text.strip()
        assert notes, f"slide {n} has no speaker notes"
        if n in NO_NUMBER:
            continue
        assert re.search(rf"^{n}$", _all_text(slide), re.M), (
            f"slide {n} is missing its slide-number textbox"
        )


def test_fonts_and_colors_stay_on_palette() -> None:
    prs = _prs()
    bad_fonts, bad_colors = set(), set()

    def inspect(text_frame):
        for para in text_frame.paragraphs:
            for run in para.runs:
                if run.font.name not in ALLOWED_FONTS:
                    bad_fonts.add(run.font.name)
                try:
                    rgb = run.font.color.rgb
                except (AttributeError, TypeError):
                    continue
                if rgb is not None and str(rgb) not in ALLOWED_COLORS:
                    bad_colors.add(str(rgb))

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                inspect(shape.text_frame)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        inspect(cell.text_frame)

    assert not bad_fonts, f"off-system fonts: {sorted(bad_fonts)}"
    assert not bad_colors, f"off-palette colours: {sorted(bad_colors)}"


def test_content_matches_source_markdown() -> None:
    """Key claims from review2_slides.md must survive into the deck."""
    prs = _prs()
    everything = "\n".join(_all_text(s) for s in prs.slides)
    everything += "\n".join(
        s.notes_slide.notes_text_frame.text for s in prs.slides
    )

    must_appear = [
        "found means fixed",          # the repositioning evidence
        "40.95",                      # Paper 2 headline number
        "77.55",                      # Paper 3 headline number
        "Copilot",                    # the competitor named openly
        "DefectDojo",
        "OS Command Injection",       # the hardest scoped CWE
        "Function preserved",         # gate 3
        "FIX_REJECTED",
        "PyGoat",
        "Cohen",                      # the agreement metric
        "Backslash",                  # the competitor on the prompting axis
        "Effective Safety Rate",      # SecCoderX's metric, our headline number
    ]
    missing = [m for m in must_appear if m not in everything]
    assert not missing, f"content dropped in transcription: {missing}"

    assert SLIDES_MD.exists(), "source markdown is missing"


def test_paper_attributions_are_correct() -> None:
    """Each new component credits the right paper.

    Easy to transpose and embarrassing in front of a panel that assigned the
    papers. Paper 1 = explained feedback (Explainer), Paper 2 = prompt syntax
    (Linter), Paper 3 = prompting techniques (Strategy registry).
    """
    prs = _prs()
    arch = _all_text(prs.slides[9])  # slide 10 — architecture diagram

    pairs = [
        ("EXPLAINER", "Paper 1"),
        ("PROMPT BUILDER", "Paper 3"),
        ("PROMPT SECURITY LINTER", "Paper 2"),
    ]
    for component, paper in pairs:
        line = next((l for l in arch.splitlines() if component in l), None)
        assert line is not None, f"architecture diagram is missing {component}"
        assert paper in line, (
            f"{component} is credited to the wrong paper: {line.strip()!r} "
            f"(expected {paper})"
        )


def test_no_emoji_in_slide_text() -> None:
    """Emoji render inconsistently in PowerPoint; ticks must be plain glyphs."""
    prs = _prs()
    offenders = []
    for n, slide in enumerate(prs.slides, start=1):
        for ch in _all_text(slide):
            if ord(ch) > 0x2800:  # past the symbol ranges we intentionally use
                offenders.append(f"slide {n}: U+{ord(ch):04X} {ch!r}")
    assert not offenders, "emoji / exotic glyphs found:\n  " + "\n  ".join(offenders[:10])


if __name__ == "__main__":
    failures = 0
    for name, fn in sorted(
        (n, f) for n, f in globals().items() if n.startswith("test_") and callable(f)
    ):
        try:
            fn()
            print(f"  PASS  {name}")
        except AssertionError as exc:
            failures += 1
            print(f"  FAIL  {name}\n        {exc}")
        except Exception as exc:  # noqa: BLE001
            failures += 1
            print(f"  ERROR {name}\n        {type(exc).__name__}: {exc}")
    print()
    print(f"  {'FAILED' if failures else 'ALL PASSED'}  ({failures} failure(s))")
    raise SystemExit(1 if failures else 0)
