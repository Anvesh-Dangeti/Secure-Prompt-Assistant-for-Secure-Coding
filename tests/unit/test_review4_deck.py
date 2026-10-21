"""Structural and factual checks on the Review 4 deck.

Builds the deck into a temporary file (so the test never depends on a stale
copy, and never fights PowerPoint's file lock), then checks layout, house
style, and -- most importantly -- that every result shown on a slide is the
recorded, test-asserted value from evaluation/recorded_results.py.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT / "scripts"))

from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

import build_review2_deck as style  # noqa: E402
import build_review4_deck as deck  # noqa: E402
from evaluation.recorded_results import LAB_REPLAYS, LINTER, TESTS_AT_REVIEW3, TESTS_AT_REVIEW4, VERIFICATION  # noqa: E402

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
NO_NUMBER = {1, 16}


@pytest.fixture(scope="module")
def prs(tmp_path_factory):  # noqa: ANN001, ANN201
    return deck.build(tmp_path_factory.mktemp("deck") / "review4.pptx")


def _text(slide) -> str:  # noqa: ANN001
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out.append(shape.text_frame.text)
        if shape.has_table:
            out.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(out)


def _rows(slide) -> list[list[str]]:  # noqa: ANN001
    table = next(s.table for s in slide.shapes if s.has_table)
    return [[c.text for c in row.cells] for row in list(table.rows)[1:]]


def test_slide_order(prs) -> None:  # noqa: ANN001
    expected = [
        "Building a Prompt Assistant", "Promised at Review 3, delivered at Review 4", "Your two questions",
        "How a developer uses it", "Nine strategies, one security envelope",
        "The linter catches what shortening removes", "Repository code is data, never instructions",
        "Four gates decide", "Eight commits, every verdict", "The scanners went quiet. The fix was rejected.",
        "Lab mode: six recorded model answers", "The dashboard: the same loop",
        "A dashboard that displays untrusted text", "Seven defects testing found", "What comes next",
        "Thank You", "Prepared answers", "Prepared answers",
    ]
    assert len(prs.slides) == len(expected)
    for i, fragment in enumerate(expected):
        assert fragment in _text(prs.slides[i]), f"slide {i + 1} missing {fragment!r}"


def test_every_shape_is_on_canvas(prs) -> None:  # noqa: ANN001
    for n, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.left is None:
                continue
            assert shape.left >= 0 and shape.top >= 0, f"slide {n}"
            assert shape.left + shape.width <= SLIDE_W and shape.top + shape.height <= SLIDE_H, f"slide {n}"


def test_tables_are_not_covered_by_later_shapes(prs) -> None:  # noqa: ANN001
    for n, slide in enumerate(prs.slides, start=1):
        shapes = list(slide.shapes)
        for i, shape in enumerate(shapes):
            if not shape.has_table:
                continue
            top, bottom = shape.top, shape.top + shape.height
            for later in shapes[i + 1:]:
                if later.shape_type not in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.PICTURE) or later.left is None:
                    continue
                overlap = (min(bottom, later.top + later.height) - max(top, later.top)) / 914400
                horizontal = later.left < shape.left + shape.width and later.left + later.width > shape.left
                assert not (horizontal and overlap > 0.02), f"slide {n}: table covered by {overlap:.2f}in"


def test_numbers_notes_fonts_and_palette(prs) -> None:  # noqa: ANN001
    fonts = {style.DISPLAY, style.BODY, style.MONO, None}
    colours = {str(c) for c in style.PALETTE}
    for n, slide in enumerate(prs.slides, start=1):
        assert slide.notes_slide.notes_text_frame.text.strip(), f"slide {n} has no speaker notes"
        if n not in NO_NUMBER:
            assert re.search(rf"^{n}$", _text(slide), re.M), f"slide {n} lacks its number"
        frames = [s.text_frame for s in slide.shapes if s.has_text_frame]
        frames += [c.text_frame for s in slide.shapes if s.has_table for r in s.table.rows for c in r.cells]
        for tf in frames:
            for para in tf.paragraphs:
                for run in para.runs:
                    assert run.font.name in fonts, f"slide {n}: font {run.font.name}"
                    rgb = run.font.color.rgb if run.font.color and run.font.color.type else None
                    assert rgb is None or str(rgb) in colours, f"slide {n}: colour {rgb}"


def test_verdict_table_shows_the_asserted_results(prs) -> None:  # noqa: ANN001
    rows = _rows(prs.slides[8])
    assert len(rows) == len(VERIFICATION)
    for cells, (label, (verdict, gate, findings)) in zip(rows, VERIFICATION.items()):
        assert cells[0] == label.split("-")[0]
        assert cells[2] == str(findings) and cells[3] == verdict, label
        decided = cells[4] == "all gates" if gate is None else cells[4].startswith(f"gate {gate}")
        assert decided, label


def test_replay_table_shows_the_asserted_results(prs) -> None:  # noqa: ANN001
    rows = _rows(prs.slides[10])
    assert [r[0] for r in rows] == list(LAB_REPLAYS)
    for cells, (rule, verdict, gate) in zip(rows, LAB_REPLAYS.values()):
        assert cells[1] == rule and cells[3] == verdict
        assert cells[4] == ("passed all four" if gate is None else f"gate {gate}")


def test_strategy_table_and_linter_numbers(prs) -> None:  # noqa: ANN001
    rows = _rows(prs.slides[4])
    n = LINTER["findings"]
    for cells in rows:
        passed, score = LINTER["by_strategy"][cells[0]]
        assert cells[-2:] == [f"{passed} / {n}", f"{score:.2f}"], cells[0]
    text = _text(prs.slides[5])
    lo, hi = LINTER["s7_tokens"]
    assert f"S7 passes {LINTER['by_strategy']['S7'][0]}/{n}" in text
    assert f"fix-its repair S0 on {LINTER['fixits_repair_s0']}/{n}" in text and f"{lo}–{hi} tokens" in text


def test_screenshots_are_embedded(prs) -> None:  # noqa: ANN001
    for name in deck.SCREENSHOTS:
        assert (deck.IMG / name).is_file(), f"missing docs/img/{name}"
    pictures = {n: sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in slide.shapes)
                for n, slide in enumerate(prs.slides, start=1)}
    assert pictures[6] == 1 and pictures[10] == 1 and pictures[12] == 2
    assert sum(pictures.values()) == len(deck.SCREENSHOTS)
    assert "screenshot missing" not in "\n".join(_text(s) for s in prs.slides)


def test_key_claims_present(prs) -> None:  # noqa: ANN001
    everything = "\n".join(_text(s) + s.notes_slide.notes_text_frame.text for s in prs.slides)
    for claim in [f"{TESTS_AT_REVIEW4} automated tests", f"{TESTS_AT_REVIEW3} at Review 3", "no plugin", "UNVERIFIED",
                  "FIX_REJECTED", "NOT_APPLICABLE", "sqli_wrong_behaviour", "Content-Security-Policy",
                  "18 November", "Copilot", "safety rate"]:
        assert claim in everything, f"missing claim: {claim}"


def test_no_emoji(prs) -> None:  # noqa: ANN001
    for n, slide in enumerate(prs.slides, start=1):
        bad = [ch for ch in _text(slide) if ord(ch) > 0x2800]
        assert not bad, f"slide {n}: {bad[:5]}"
