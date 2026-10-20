#!/usr/bin/env python3
"""
Build the Review 4 presentation deck.

The design system (palette, fonts, grid, helpers) is imported from
build_review2_deck.py so the series stays visually identical. Content is
mirrored in docs/review4_slides.md. Every result number is imported from
evaluation/recorded_results.py -- the same constants the tests assert against
real runs -- so no slide can show a number the code does not actually produce.
The dashboard screenshots live in docs/img/ and were taken from the running
dashboard on the eight-commit demo database (scripts/demo_verify.py).

Usage:
    python scripts/build_review4_deck.py
    python scripts/build_review4_deck.py --out path/to/deck.pptx
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(ROOT))

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

import build_review2_deck as base  # noqa: E402
from build_review2_deck import (  # noqa: E402
    BORDER, CONTENT_W, DEEP1, DEEP2, DISPLAY, MARGIN_X, NAVY, PALE, SLATE, TEAL, WHITE, YES,
    add_slide, badge, callout, data_table, headline, mono_block, notes, section_label, slide_number, textbox,
    write,
)
from app.ai.strategies import STRATEGIES  # noqa: E402
from evaluation.recorded_results import (  # noqa: E402
    FIXTURE_BEHAVIOUR_TESTS, LAB_REPLAYS, LINTER, TESTS, TESTS_AT_REVIEW3, VERIFICATION,
)

IMG = ROOT / "docs" / "img"
SCREENSHOTS = ("review4_overview.png", "review4_prompt.png", "review4_lint.png", "review4_verification.png")

VERIFICATION_STORIES = {
    "v2-stub-ping": "ping_host emptied to `pass` — every scanner goes quiet",
    "v3-fix-ping": "ping_host fixed: argument list, no shell",
    "v4-churn-sql": "the SQL line edited, but still concatenated",
    "v5-fix-sql-new-route": "SQL parameterised — and a new os.system() route in the same commit",
    "v6-remove-route": "the os.system() route removed",
    "v7-token-gitleaks-down": "token moved to an environment variable — while Gitleaks is down",
    "v8-gitleaks-back": "same code, Gitleaks available again",
}
REPLAY_STORIES = {
    "sqli_good": "parameterised query",
    "sqli_stub": "body replaced by raise NotImplementedError",
    "sqli_drops_route": "query fixed, but the @app.route decorator lost",
    "sqli_wrong_behaviour": "scanner-clean and structurally intact — but queries the wrong column",
    "cmd_quoted_shell": "shlex.quote(), but still shell=True",
    "cmd_good": "argument list, no shell",
}


def _circles(s) -> None:
    for x, y, d, col in ((9.40, 1.10, 3.60, DEEP1), (10.00, 1.70, 2.40, DEEP2), (10.55, 2.25, 1.30, TEAL)):
        shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        shp.fill.solid()
        shp.fill.fore_color.rgb = col
        shp.line.fill.background()
        shp.shadow.inherit = False


def _white_box(slide, x, y, w, h):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = WHITE
    shp.line.fill.background()
    shp.shadow.inherit = False
    shp.adjustments[0] = 0.04
    return shp


def picture(slide, name, x, y, w, *, crop=(0.0, 0.0, 0.0, 0.0), caption=None):
    """A dashboard screenshot with a hairline border, optionally cropped (left, top, right, bottom fractions).

    A missing file shows a labelled frame instead, so the deck still builds (the deck test fails on it).
    """
    path = IMG / name
    left, top, right, bottom = crop
    h = w * 9 / 16 * (1 - top - bottom) / (1 - left - right)
    if path.exists():
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
        px_w, px_h = pic.image.size
        h = w * (px_h * (1 - top - bottom)) / (px_w * (1 - left - right))
        pic.height = Inches(h)
        pic.line.color.rgb = BORDER
        pic.line.width = Pt(0.75)
    else:
        _white_box(slide, x, y, w, h)
        tf = textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.MIDDLE)
        write(tf, f"screenshot missing: docs/img/{name}", size=14, color=SLATE, align=PP_ALIGN.CENTER)
    if caption:
        tf = textbox(slide, x, y + h + 0.08, w, 0.30)
        write(tf, caption, size=11, color=SLATE)
    return h


def _panel(slide, x, w, label, question, lines, *, y=1.95, h=4.05):
    _white_box(slide, x, y, w, h)
    tf = textbox(slide, x + 0.30, y + 0.22, w - 0.60, 0.30)
    write(tf, label.upper(), font=DISPLAY, size=12, bold=True, color=TEAL)
    tf = textbox(slide, x + 0.30, y + 0.55, w - 0.60, 0.75)
    write(tf, question, font=DISPLAY, size=15, bold=True, color=NAVY, line_spacing=1.05)
    tf = textbox(slide, x + 0.30, y + 1.38, w - 0.60, h - 1.55)
    for line in lines:
        write(tf, line, size=12, color=SLATE, space_after=7, line_spacing=1.1)


def _mini_box(slide, col, y, title, body, *, h=0.90):
    x = base.CARD_X[col]
    _white_box(slide, x, y, base.CARD_W, h)
    tf = textbox(slide, x + 0.22, y + 0.10, base.CARD_W - 0.44, 0.30)
    write(tf, title, font=DISPLAY, size=13, bold=True, color=NAVY)
    tf = textbox(slide, x + 0.22, y + 0.40, base.CARD_W - 0.44, h - 0.45)
    write(tf, body, size=11, color=SLATE, line_spacing=1.05)


def _note_row(slide, index, num, title, body, *, x=9.45, y0=1.95, pitch=1.35, w=3.25):
    y = y0 + index * pitch
    badge(slide, x, y, 0.55, num, shape=MSO_SHAPE.ROUNDED_RECTANGLE, label_size=12)
    tf = textbox(slide, x + 0.72, y, w - 0.72, 0.35)
    write(tf, title, font=DISPLAY, size=15, bold=True, color=NAVY)
    tf = textbox(slide, x + 0.72, y + 0.38, w - 0.72, 1.05)
    write(tf, body, size=11.5, color=SLATE, line_spacing=1.1)


def _gate_cell(verdict: str, gate: str | None) -> str:
    if gate is None:
        return "all gates"
    return f"gate {gate} — no evidence" if verdict == "UNVERIFIED" else f"gate {gate}"


# ==========================================================================
# SLIDES
# ==========================================================================

def slide_01_title(prs):
    s = add_slide(prs, NAVY)
    _circles(s)
    tf = textbox(s, 0.70, 2.15, 8.20, 0.40)
    write(tf, "SECURE SOFTWARE ENGINEERING · REVIEW 4", font=DISPLAY, size=13, bold=True, color=TEAL)
    tf = textbox(s, 0.70, 2.60, 8.60, 1.90)
    write(tf, "Building a Prompt Assistant", font=DISPLAY, size=40, bold=True, color=WHITE)
    write(tf, "for Secure Coding", font=DISPLAY, size=40, bold=True, color=WHITE)
    tf = textbox(s, 0.70, 4.55, 8.00, 0.90)
    write(tf, "Phases 5–7 built: prompt, lint, verify — and a dashboard to drive it",
          size=16, color=PALE, line_spacing=1.2)
    for name, roll, x in (("Student 1 : Anvesh Dangeti", "Roll no: EM26MTECH14009", 0.70),
                          ("Student 2 : Sai Prakash Nayak", "Roll no: EM26MTECH11001", 3.57),
                          ("Student 3 : Dhananjay Kumar", "Roll no: CS26MTECH13007", 6.44)):
        tf = textbox(s, x, 6.34, 2.60, 0.46)
        write(tf, name, size=12, bold=True, color=PALE)
        tf = textbox(s, x, 6.74, 2.60, 0.38)
        write(tf, roll, size=12, bold=True, color=PALE)
    notes(s, """
Review 4 - Saturday 24 October 2026. Fifteen minutes.

Frame: Review 3 showed scanning and tracking. This review shows the contribution
working end to end - prompt, lint, the developer's own assistant, rescan, gates -
and answers the two questions the panel asked at Review 3.

Load-bearing slides: 2, 3, 4, 9, 10. Cuttable if short on time: 5, 7, 11, 14.
Have a terminal open with the demo already run once (pip-audit needs internet).
""")


def slide_02_delivered(prs):
    s = add_slide(prs)
    section_label(s, "Since Review 3")
    headline(s, "Promised at Review 3, delivered at Review 4")
    data_table(
        s,
        ["Promised for Review 4 (Review 3, slide 12)", "Status"],
        [
            ["Phase 5 AI layer — explainer, prompt builder × S0–S8, prompt linter, injection defences",
             f"{YES}  done"],
            ["Phase 6 Verification — rescan, gates 1 · 2 · 3, VERIFIED_RESOLVED / FIX_REJECTED, history",
             f"{YES}  done, four gates"],
            ["Phase 7 Dashboard — promised for Review 5", f"{YES}  done, two weeks early"],
            ["The panel's two questions from Review 3", f"{YES}  answered — next slide"],
            ["The 12 original requirements — #11's measured evaluation is Phase 8", f"{YES}  all 12 built"],
        ],
        [9.10, 3.00],
        row_h=0.62, body_pt=12.5, bold_rows=(2,),
    )
    callout(s, f"{TESTS} automated tests, up from {TESTS_AT_REVIEW3} at Review 3 — and every result in this "
               "deck is a recorded value that a test asserts.", emphasis=True)
    slide_number(s, 2)
    notes(s, f"""
Say: "At Review 3 we committed to Phases 5 and 6 for today. Both are done, and the
dashboard, which we had scheduled for Review 5, is done as well."

All twelve original requirements now have working code. Requirement 11 - testing
on representative OWASP Top 10 cases and vulnerable apps - is met by the fixture
(10 CWEs, 7 OWASP categories) and vulpy; its measured evaluation is Phase 8.

{TESTS} automated tests ({TESTS_AT_REVIEW3} at Review 3). The repository is private, with
per-phase commits.
""")


def slide_03_questions(prs):
    s = add_slide(prs)
    section_label(s, "Your questions from Review 3")
    headline(s, "Your two questions, answered")
    _panel(s, 0.60, 5.95, "You asked", "“In what mode will a developer use this — an IDE plugin, or something else?”", [
        "Assistant-agnostic: no plugin. It runs beside whatever assistant the developer already uses — "
        "Copilot, Cursor, Codex.",
        "Developer mode: we scan, explain and build a linted prompt; their assistant writes the fix; we "
        "rescan their commit and the gates decide.",
        "No model of ours, no tokens from us. Our local model is only a lab instrument for the strategy "
        "experiments.",
        "CLI and dashboard today; an editor extension stays a Could, not a Must.",
    ])
    _panel(s, 6.75, 5.95, "You asked", "“Are we overcommitting? Can the project be finished by 18 November?”", [
        "Scope is tiered. Must: S0–S8 + linter, injection defences, gated verification, dashboard, "
        "evaluation on our fixture, hardening, documents.",
        "Should and Could are named in advance, with a cut rule: a phase more than 2 days late drops the "
        "next Should item — never tests or hardening.",
        "Phases 5–7 finished on 20 October; Phase 7 was planned for 25–31 October.",
        "What is left — evaluation, the web interface, hardening, documents — is dated on slide 15.",
    ])
    callout(s, "Every design decision traces to a paper or a measurement — and we reversed two when the "
               "evidence said so (Bandit restored; patch-based apply dropped).", y=6.25, h=0.62)
    slide_number(s, 3)
    notes(s, """
Question 1. The tool is a companion, not an extension of one assistant. A plugin
would tie us to one IDE and one assistant; the requirement says "Copilot or
another coding assistant". The developer copies our prompt into the assistant
they already pay for. Slide 4 shows the exact loop.

Question 2. The dated plan is on slide 15. Must is what the brief requires; Should
and Could were written down before we started, so cutting them is a plan, not a
failure. The two reversals: Bandit was dropped at Review 0 and restored after
Paper 3 used it as its primary detector; git-worktree patches were dropped because
small models produce diffs that do not apply.
""")


def slide_04_developer_loop(prs):
    s = add_slide(prs)
    section_label(s, "Developer mode")
    headline(s, "How a developer uses it — no plugin needed")
    mono_block(s, """
# 1  scan the checkout: four scanners, every finding tracked by identity
python -m app.cli scan-zip ./myapp --project myapp
python -m app.cli findings --project myapp

# 2  build the prompt for one finding (S7 by default, linted before it is shown)
python -m app.cli prompt 5 --source ./myapp --out prompt.txt

# 3  edited or shortened it?  lint it again; --fix puts back what was cut
python -m app.cli lint-prompt prompt.txt --finding 5 --source ./myapp --fix

# 4  paste it into Copilot / Cursor / Codex, apply the change, commit

# 5  rescan, and let the gates decide
python -m app.cli verify 5 --source ./myapp          <-- VERIFIED_RESOLVED | FIX_REJECTED | UNVERIFIED
""", y=1.90, h=4.25, size=13)
    callout(s, "The dashboard drives the same services: pick a finding, copy the prompt, lint your edit, rescan "
               "and verify. Developer mode never executes repository code.", y=6.35, h=0.62)
    slide_number(s, 4)
    notes(s, """
RUN THE LOOP HERE if there is a laptop (presenter script, demo step 2).

Step 2: the explanation comes from a template - no model is called and no tokens
are spent by us. Step 3 is the part no other tool does: developers shorten
prompts, and shortening is exactly what removes the security constraints. The
linter checks the developer's own text, not just ours.

Step 5 compares the new scan with the latest scan in which the finding was still
present, so a rejected fix can be corrected in a follow-up commit and verified.
""")


def slide_05_strategies(prs):
    s = add_slide(prs)
    section_label(s, "Prompt strategies")
    headline(s, "Nine strategies, one security envelope")
    rows = []
    for sid, strategy in STRATEGIES.items():
        passed, score = LINTER["by_strategy"][sid]
        rows.append([sid, strategy.name, strategy.summary, strategy.origin,
                     f"{passed} / {LINTER['findings']}", f"{score:.2f}"])
    data_table(s, ["ID", "Strategy", "What it adds", "Source", "Passes lint", "Constituents"], rows,
               [0.60, 1.95, 4.00, 3.30, 1.05, 1.20], row_h=0.40, header_pt=11, body_pt=10.5,
               bold_rows=(list(STRATEGIES).index("S7"),))
    callout(s, "Delimiters, the inert-data instruction, the scope limit and the output contract are identical in "
               "all nine — only the guidance varies. This scores the prompts, not the fixes: which strategy "
               "yields verified fixes is Phase 8 (E2).", y=6.20, h=0.70, size=12)
    slide_number(s, 5)
    notes(s, f"""
Measured on the {LINTER['findings']} real findings the scanners report on our fixture, with no model involved.

S0 is "direct feedback": the raw scanner message, no CWE, no guidance - the
negative control. S1-S6 are Paper 3's strategies; S7 is ours; S8 is the
rules-file technique used in industry.

Only S7 passes on every finding, by construction: it is the linted template, and
the builder refuses to emit an S7 prompt that fails. S2 and S8 score highest of the
rest because they name the CWE's mitigation.

Cuttable if short on time.
""")


def slide_06_linter(prs):
    s = add_slide(prs)
    section_label(s, "Prompt security linter")
    headline(s, "The linter catches what shortening removes")
    # Cropped to the verdict and the constituent table, so the text is legible at slide size.
    picture(s, "review4_lint.png", MARGIN_X, 1.85, CONTENT_W, crop=(0.232, 0.398, 0.045, 0.268),
            caption="Dashboard: a shortened prompt — “Fix the path traversal in download_file in app.py” plus "
                    "the code — names the weakness and the place, and loses everything else.")
    _mini_box(s, 0, 5.45, "A1–A7  constituents",
              "the input at risk, a guard, a directive, a scope limit, the opening — and what an edit removed")
    _mini_box(s, 1, 5.45, "B1–B5  prompt security",
              "no secrets, intact delimiters, the inert-data instruction, data minimisation, injection-like code")
    _mini_box(s, 2, 5.45, "C1  budget  ·  fix-its",
              "a token budget, and every error comes with the sentence that repairs it")
    lo, hi = LINTER["s7_tokens"]
    s0_lo, s0_hi = LINTER["s0_score_range"]
    n = LINTER["findings"]
    callout(s, f"On the {n} real findings:  S7 passes {LINTER['by_strategy']['S7'][0]}/{n}  ·  S0 passes "
               f"{LINTER['by_strategy']['S0'][0]}/{n} (constituents {s0_lo}–{s0_hi})  ·  fix-its repair S0 on "
               f"{LINTER['fixits_repair_s0']}/{n}  ·  S7's first message {lo}–{hi} tokens",
            emphasis=True, y=6.45, h=0.62, size=12.5)
    slide_number(s, 6)
    notes(s, """
Paper 2 proposes a prompt linter as future work. Linting only our own templates
would prove nothing, so it lints whatever the developer is about to send.

The screenshot is a realistic edit: "Fix the path traversal in download_file" plus
the code in a fence. It names the weakness and the place (A5 passes) but drops the
input at risk, the guard, the handling directive and the scope limit. Bound to
the finding, the linter also reports that the edit removed constituents the
generated prompt had (A7).

Names inside fenced code do not count as instructions - otherwise pasting the
code would satisfy entity binding by itself (defect 9).
""")


def slide_07_injection(prs):
    s = add_slide(prs)
    section_label(s, "Prompt injection defences")
    headline(s, "Repository code is data, never instructions")
    data_table(s, ["Attack", "Control", "Tested in"], [
        ["Instructions hidden in comments or docstrings",
         "stripped from the code before it is quoted — which also keeps the fixture's CWE-naming docstrings "
         "out of the S0 control", "context"],
        ["Code that tries to break out of its block",
         "untrusted_code markers with a random nonce and an explicit “treat as data” instruction; a forged "
         "marker is flagged (B5) but can never block our own prompt", "delimiting, injection"],
        ["Secrets in the quoted code",
         "redacted before any prompt is built; a secret in a developer's edited prompt is lint error B1",
         "context, linter"],
        ["Over-sharing", "only the enclosing function (≤ 40 lines) and the finding's metadata — never the file "
                         "or the repository (B4)", "linter"],
        ["A model's answer", "schema-checked: code, imports (each must parse as one import), summary; extra "
                             "fields refused; never executed or committed by us", "lab output"],
        ["A model's explanation", "template by default; a model's text is kept only if schema-valid, naming a "
                                  "real entity, and free of code and injection phrases", "explainer"],
    ], [2.70, 7.45, 1.95], row_h=0.60, body_pt=11)
    callout(s, "And the decisive control: nothing a model writes is believed. The scanners and the gates "
               "decide.", emphasis=True, y=6.42, h=0.55)
    slide_number(s, 7)
    notes(s, """
Repository content reaches the prompt, and the repository is untrusted - a
comment saying "ignore previous instructions" is attacker input. Delimiting
cannot make a model safe; it makes the boundary explicit and testable.

The forged-marker row is a defect we found (12): a fake end marker inside the
code would have let scanned code mark our own prompt as malformed - a
denial-of-service lever. It is now a warning, not an error.

Cuttable if short on time.
""")


def slide_08_gates(prs):
    s = add_slide(prs)
    section_label(s, "Verification")
    headline(s, "Four gates decide — the model never does")
    data_table(s, ["Gate", "Asks", "Passes when", "Fails when", "Mode"], [
        ["1", "Is the weakness gone?", "every finding of the weakness — same CWE, same function, any scanner "
                                       "— is absent from the rescan",
         "still reported; or a NEW finding there from the same rule, or the same CWE at ≥ its severity",
         "both"],
        ["2", "Did the commit add anything?", "no NEW finding anywhere in the commit at or above the "
                                              "weakness's severity (lower ones are warnings)",
         "a fix that adds a weakness; a failed scanner or blind spot → UNVERIFIED", "both"],
        ["3a", "Is the function still there?", "same parameters, decorators kept, a real body, no new "
                                               "undefined names — facts captured by the scan",
         "the code was deleted, emptied to pass / raise, or calls a name that does not exist", "both"],
        ["3b", "Does it still behave?", f"{FIXTURE_BEHAVIOUR_TESTS} behaviour tests pass in an isolated "
                                        "interpreter, run only after 1–3a pass",
         "the fix changes what the function returns", "lab only"],
    ], [0.65, 2.30, 4.55, 3.50, 1.10], row_h=0.82, body_pt=11)
    callout(s, "Every gate passes → VERIFIED_RESOLVED.   A gate fails → FIX_REJECTED, reason recorded.   "
               "Evidence missing → UNVERIFIED.   Never a silent pass.", emphasis=True, y=6.40, h=0.62)
    slide_number(s, 8)
    notes(s, """
Gate 1's target is the weakness, not one finding: Bandit B608 and both Semgrep SQL
rules on get_user are one target of three findings, and all three must go.

The churn rule: when a commit edits a vulnerable line, the scanner reports the
same weakness under a new fingerprint. A NEW finding in the same function from
the same rule - or the same CWE at or above the weakness's severity - counts as
"still present". Lower-severity advisories become gate-2 warnings (defect 13).

Developer mode never runs repository code, so 3b is reported as NOT_APPLICABLE -
not as passed.
""")


def slide_09_verdicts(prs):
    s = add_slide(prs)
    section_label(s, "Verification — live demonstration")
    headline(s, "Eight commits, every verdict")
    rows = []
    for label, (verdict, gate, findings) in VERIFICATION.items():
        rows.append([label.split("-")[0], VERIFICATION_STORIES[label], str(findings), verdict,
                     _gate_cell(verdict, gate)])
    data_table(s, ["", "What the developer committed (after the v1 baseline)", "Findings", "Verdict",
                   "Decided by"], rows, [0.55, 6.25, 1.10, 2.60, 1.60], row_h=0.46, header_pt=11, body_pt=11,
               bold_rows=(0,))
    callout(s, "v2: the lifecycle says RESOLVED — every scanner stopped reporting it. Verification says "
               "FIX_REJECTED: the function was emptied.", emphasis=True, y=5.80, h=0.62)
    callout(s, "Run it:   python scripts/demo_verify.py   —   real scanners, real upload path, asserted by an "
               "integration test.", y=6.52, h=0.48, size=12)
    slide_number(s, 9)
    notes(s, """
RUN THE DEMO HERE if there is a laptop (about 50 seconds; presenter script step 3).

v3 is verified against v1 - the latest scan in which the weakness was present -
not against v2. v5 fixes the SQL injection but adds a new HIGH finding in the same
commit: gate 2 rejects it, and v6 removes the route and the same fix verifies.

v7: Gitleaks could not run, so its view of the secret cannot be checked - the
verdict is UNVERIFIED, never VERIFIED. v8, with Gitleaks back, verifies.
""")


def slide_10_quiet_but_rejected(prs):
    s = add_slide(prs)
    section_label(s, "The headline result")
    headline(s, "The scanners went quiet. The fix was rejected.")
    # Cropped to the v2 verification card, so the gate table is legible at slide size.
    picture(s, "review4_verification.png", MARGIN_X, 1.95, 8.60, crop=(0.232, 0.465, 0.045, 0.0),
            caption="Dashboard → Verification, finding #5 (Bandit B602 in ping_host): the v2 commit.")
    _note_row(s, 0, "01", "Gate 1 passes", "4 of 4 CWE-78 findings are gone from ping_host: no scanner reports "
                                          "the weakness any more.")
    _note_row(s, 1, "02", "Gate 3a fails", "the body was reduced to a stub. Verdict FIX_REJECTED, with the "
                                          "reason recorded.")
    _note_row(s, 2, "03", "The next commit", "a real fix — argument list, no shell — passes every gate and is "
                                            "VERIFIED_RESOLVED.")
    callout(s, "A scanner-only check would accept v2: deleting vulnerable code deletes the finding too.",
            emphasis=True, y=6.20, h=0.62)
    slide_number(s, 10)
    notes(s, """
This is the point of the project in one screen. A scanner-only check - "is the
finding gone?" - would accept v2, because deleting vulnerable code also deletes
the finding. Paper 3 documents exactly these fake fixes.

Lower card: v2. Gate 1 PASS, gate 2 PASS, gate 3a FAIL. Upper card: v3, every
gate passes; the two LOW Bandit advisories on the new subprocess call are
reported as gate-2 warnings, not failures.
""")


def slide_11_lab(prs):
    s = add_slide(prs)
    section_label(s, "Lab mode")
    headline(s, "Lab mode: six recorded model answers, four gates")
    rows = []
    for name, (rule, verdict, gate) in LAB_REPLAYS.items():
        rows.append([name, rule, REPLAY_STORIES[name], verdict, "passed all four" if gate is None else f"gate {gate}"])
    data_table(s, ["Recorded answer", "Finding", "What the model's fix did", "Verdict", "Stopped at"], rows,
               [2.45, 0.95, 4.90, 2.30, 1.50], row_h=0.50, header_pt=11, body_pt=11,
               bold_rows=(list(LAB_REPLAYS).index("sqli_wrong_behaviour"),))
    callout(s, "sqli_wrong_behaviour passes gates 1, 2 and 3a — only the behaviour tests catch it. That is why lab "
               "mode runs them, and why developer mode reports 3b as not applicable rather than passed.",
            emphasis=True, y=5.60, h=0.70, size=12)
    callout(s, "Recorded answers replayed through the real pipeline — real scanners, real tests, scratch copy. "
               "Live local-model runs are Phase 8.", y=6.42, h=0.52, size=12)
    slide_number(s, 11)
    notes(s, f"""
Lab mode is for the experiments: the model's answer replaces the function in a
scratch copy of the fixture, both copies are rescanned, and gates 1, 2, 3a and
then 3b run. {FIXTURE_BEHAVIOUR_TESTS} behaviour tests pass on the unmodified fixture; process calls are faked
so no test can run a real command.

cmd_quoted_shell is a classic AI fix: shlex.quote() looks right, but shell=True is
still there and Bandit still reports it - gate 1 rejects it.

Cuttable if short on time.
""")


def slide_12_dashboard(prs):
    s = add_slide(prs)
    section_label(s, "Dashboard")
    headline(s, "The dashboard: the same loop, on one screen")
    picture(s, "review4_overview.png", 0.60, 1.95, 5.90,
            caption="Overview — severity, lifecycle, latest verdicts, OWASP spread, scan trend")
    picture(s, "review4_prompt.png", 6.80, 1.95, 5.90,
            caption="Prompt & lint — the S7 prompt, ready to copy into any assistant")
    callout(s, "Five pages: Overview · Findings (suppress with a reason) · Prompt & lint · Verification · Scan "
               "history. In-process over the CLI's services — no lab mode, no model, nothing executed.",
            y=5.95, h=0.70)
    slide_number(s, 12)
    notes(s, """
Built on the same services as the CLI, so the dashboard and lint-prompt give
identical lint reports (asserted by a test).

Every chart has its numbers in a table on the same page. The prompt is shown as
code with a copy button: the developer pastes it into their own assistant.

The screenshots are from the eight-commit demonstration database.
""")


def slide_13_dashboard_security(prs):
    s = add_slide(prs)
    section_label(s, "Securing the dashboard")
    headline(s, "A dashboard that displays untrusted text")
    data_table(s, ["Threat", "Control", "Proven by"], [
        ["Reached from the network", "binds 127.0.0.1 and refuses to render anything — even the login — if "
                                     "configured otherwise", "test + netstat"],
        ["Used without the key", "the API's key; constant-time comparison; no key configured means closed; a 1 s "
                                 "delay and a lockout after 5 wrong keys", "tests"],
        ["Scanned text read as markup", "file names, rule ids and messages are attacker-controlled: shown only as "
                                        "text, code or table cells; no Markdown or HTML writers (AST-checked)",
         "AST test"],
        ["The payload test itself", "a Markdown image, <img onerror> and a javascript: link planted in a rule id, "
                                    "a message and a file name — and the test checks they are shown, as text",
         "every page"],
        ["Cross-site requests, leakage", "XSRF and CORS protection on; no static file serving; no tracebacks, no "
                                         "telemetry, no file watcher", "config test"],
        ["The checkout path field", "absolute, an existing directory, outside scratch space, optionally under one "
                                    "allowed root", "unit tests"],
    ], [2.75, 7.45, 1.90], row_h=0.58, body_pt=11)
    callout(s, "For Review 5 it becomes a FastAPI + htmx web interface: the same controls, plus a session cookie, "
               "CSRF tokens and a strict Content-Security-Policy.", emphasis=True, y=6.25, h=0.62)
    slide_number(s, 13)
    notes(s, """
A repository under scan controls its file names, and its code shapes scanner
messages. If any of that reached a Markdown renderer, an image link would make
the viewer's browser fetch an attacker's URL.

The payload test seeds the payload through a stand-in scanner whose rule id and
message carry it, and through a file literally named "![file](evil).py". It
visits every page and asserts the markers never appear in a Markdown element -
and that they do appear as literal text (defect 18 taught us that check).

Why rebuild: Streamlit gave us a working, secure dashboard quickly; its layout and
density limit how fast a developer can work. Server-rendered templates keep one
Python stack with no npm supply chain, and drop about 25 transitive packages.
""")


def slide_14_defects(prs):
    s = add_slide(prs)
    section_label(s, "What testing found")
    headline(s, "Seven defects testing found")
    data_table(s, ["Defect", "Found by", "Would have broken"], [
        ["The guidance term “salt” matched “unsalted”, which says the opposite", "unit test on real guidance",
         "a prompt asking for the wrong thing passes"],
        ["Two guidance guards failed their own lint rule", "property test over every CWE",
         "fix-its that never converge"],
        ["A forged delimiter in scanned code marked our own prompt malformed", "injection tests",
         "scanned code blocking its own fix prompt"],
        ["The churn rule rejected the correct ping fix (Bandit LOW advisories)", "the eight-commit demo",
         "a right fix reported as FIX_REJECTED"],
        ["Gate 3b imported our own app/ package instead of the fixture's app.py", "the first lab run",
         "behaviour tests run against the wrong code"],
        ["The dashboard injection test passed without ever seeding the payload", "re-reading its evidence",
         "a security test that proved nothing"],
        ["A smoke run read an empty database (shell path conversion)", "checking screens against data",
         "“every page renders” meant nothing"],
    ], [5.25, 2.75, 4.10], row_h=0.56, body_pt=11)
    callout(s, "Lesson kept: every negative security test now has a positive control — it must also show the "
               "payload is displayed, as text.")
    slide_number(s, 14)
    notes(s, """
Numbered 8 to 19 in our log, continuing from Review 3's list. The full list is in
the repository's phase log.

One more, found while taking the screenshots for this deck: the OWASP bar chart
hid every other category name. Fixed, with a test.

Cuttable if short on time.
""")


def slide_15_plan(prs):
    s = add_slide(prs)
    section_label(s, "Plan to 18 November")
    headline(s, "What comes next — and how we finish by 18 Nov")
    data_table(s, ["Dates", "Phase", "Deliverable"], [
        ["25–31 Oct", "8  Evaluation", "E1 linter mutation sensitivity (no model) · E2 S0–S8 safety rate + "
                                       "tokens · E3 weak vs linted prompts"],
        ["1–7 Nov", "7b  Web interface", "FastAPI + server-rendered templates + htmx replaces Streamlit; strict "
                                         "CSP, session and CSRF"],
        ["Sat 7 Nov", "Review 5", "evaluation results and the new interface"],
        ["8–12 Nov", "9  Hardening", "security review of our own application, including the new interface"],
        ["13–16 Nov", "10  Final", "integration and regression testing; architecture and threat-model "
                                   "documents; demo"],
        ["18 Nov", "Submission", "17 November kept free as a buffer"],
    ], [1.60, 2.40, 8.10], row_h=0.52, body_pt=11.5)
    callout(s, "Questions for the panel:  (1) For E2, is one local open model across all nine strategies an "
               "acceptable scope, with a hosted model as stretch?   (2) Is a hand-pasted Copilot spot-check "
               "(10 findings, weak vs linted prompts) worth including?", emphasis=True, y=5.90, h=0.95, size=12.5)
    slide_number(s, 15)
    notes(s, """
End on the questions - hand the floor back deliberately.

Cut rule, stated at Review 3: if Phase 8 runs more than two days late, the web
interface shrinks to its core pages (Overview, Findings, Verify). Hardening and
tests are never compressed.

E1 needs no model at all: delete one constituent from an S7 prompt and check the
linter catches it. E2 runs the local model (qwen2.5-coder 7B) through lab mode on
every fixture finding, per strategy, and reports verified fixes and tokens.
""")


def slide_16_thanks(prs):
    s = add_slide(prs, NAVY)
    _circles(s)
    tf = textbox(s, 0.70, 3.05, 8.60, 1.10)
    write(tf, "Thank You", font=DISPLAY, size=40, bold=True, color=WHITE)
    tf = textbox(s, 0.70, 4.20, 8.20, 0.60)
    write(tf, "Questions welcome — demo ready on request.", size=16, color=PALE)
    notes(s, "Appendix slides 17 and 18 hold prepared answers. Jump there only if pressed.")


def slide_17_appendix_a(prs):
    s = add_slide(prs)
    section_label(s, "Appendix — anticipated questions  (1 of 2)")
    headline(s, "Prepared answers", size=28)
    base._qa_row(s, 0, "“Why not let the model decide whether its fix worked?”",
                 "A model that writes a wrong fix confidently will judge it confidently too. The scanners that "
                 "found the issue re-run and the gates decide; nothing the model writes is executed, committed or "
                 "trusted. v2 is a fix every scanner accepted that the gates still rejected.")
    base._qa_row(s, 1, "“Copilot Autofix already closes the loop — what is new here?”",
                 "We said so at Review 1. Ours is open and measurable: nine prompting strategies compared on the "
                 "same findings, a linter for the developer's own prompts, gates that include function "
                 "preservation, and finding identity tracked across commits.")
    lo, hi = LINTER["s7_tokens"]
    base._qa_row(s, 2, "“Who pays for the tokens?”",
                 f"In developer mode we call no model: the explanation is a template and the prompt is built "
                 f"locally. The developer's assistant reads about {lo}–{hi} tokens for S7's first message, and the "
                 f"linter's C1 rule warns above a budget. Lab mode runs a local model, so no per-call cost.")
    base._qa_row(s, 3, "“Why does developer mode never run the tests?”",
                 "It is the developer's repository — running its code would execute untrusted code on the scanning "
                 "machine. Gate 3a uses facts the scan already captured. Behaviour tests run only in lab mode, on "
                 "our own fixture, in an isolated interpreter with a timeout.")
    slide_number(s, 17)
    notes(s, "Appendix. Present only if asked.")


def slide_18_appendix_b(prs):
    s = add_slide(prs)
    section_label(s, "Appendix — anticipated questions  (2 of 2)")
    headline(s, "Prepared answers", size=28)
    base._qa_row(s, 0, "“Can gate 3a be fooled?”",
                 "Yes: a fix that keeps the signature but changes behaviour passes 3a. That is why lab mode has 3b, "
                 "and sqli_wrong_behaviour shows 3b catching what 3a cannot. In developer mode the verdict says 3b "
                 "was not applicable — it never claims a pass it did not observe.")
    base._qa_row(s, 1, "“A 7B model is weak — why use it at all?”",
                 "Lab mode measures differences between strategies on the same model; a model that fixed everything "
                 "would hide them. Developer mode does not depend on it — the developer's own assistant writes "
                 "the fix.")
    base._qa_row(s, 2, "“What if the developer's assistant changes other files?”",
                 "Gate 2 covers the whole commit: any new finding at or above the weakness's severity, anywhere, "
                 "rejects the fix; lower ones are reported. The prompt's scope limit (A4) asks for no unrelated "
                 "changes, and the linter checks it is still there.")
    base._qa_row(s, 3, "“Why rebuild a dashboard that works?”",
                 "Streamlit gave us a working, secure dashboard in one phase, but its layout limits how fast a "
                 "developer can work. Server-rendered templates with htmx keep one Python stack, add no npm supply "
                 "chain, and let us set the CSP, sessions and CSRF ourselves.")
    slide_number(s, 18)
    notes(s, "Appendix. Present only if asked.")


BUILDERS = [
    slide_01_title, slide_02_delivered, slide_03_questions, slide_04_developer_loop, slide_05_strategies,
    slide_06_linter, slide_07_injection, slide_08_gates, slide_09_verdicts, slide_10_quiet_but_rejected,
    slide_11_lab, slide_12_dashboard, slide_13_dashboard_security, slide_14_defects, slide_15_plan,
    slide_16_thanks, slide_17_appendix_a, slide_18_appendix_b,
]


def build(out_path: Path) -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(base.SLIDE_W)
    prs.slide_height = Inches(base.SLIDE_H)
    for builder in BUILDERS:
        builder(prs)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return prs


def main() -> int:
    default_out = ROOT.parent / "Secure_Prompt_Assistant_Review4.pptx"
    parser = argparse.ArgumentParser(description="Build the Review 4 deck")
    parser.add_argument("--out", type=Path, default=default_out)
    args = parser.parse_args()
    missing = [name for name in SCREENSHOTS if not (IMG / name).exists()]
    build(args.out)
    print(f"  built {len(BUILDERS)} slides -> {args.out}")
    if missing:
        print(f"  MISSING screenshots (placeholders used): {', '.join(missing)}")
    if base._warnings:
        print(f"\n  {len(base._warnings)} possible layout issue(s):")
        for w in base._warnings:
            print(w)
    else:
        print("  no layout warnings")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
